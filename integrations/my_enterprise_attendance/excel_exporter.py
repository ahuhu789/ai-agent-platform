import os
from datetime import datetime
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from integrations.my_enterprise_attendance.repository import get_repository


def export_monthly_attendance_to_excel(month: int, year: int, output_dir: str = "exports") -> Dict[str, Any]:
    """
    Xuất báo cáo thống kê chuyên cần tháng thành file Excel (.xlsx) hoàn chỉnh gồm:
    1. Sheet 'Tổng Quan': Số liệu tổng hợp toàn công ty.
    2. Sheet 'Chi Tiết Nhân Viên': Danh sách chi tiết chuyên cần của từng nhân sự trong tháng.
    3. Sheet 'Nhật Ký Chấm Công': Toàn bộ các bản ghi chấm công thực tế phát sinh trong tháng.
    """
    repo = get_repository()
    repo.load_data()
    
    # 1. Lấy dữ liệu thống kê tháng
    stats_data = repo.get_monthly_attendance_statistics(month, year)
    if not stats_data.get("success"):
        return stats_data

    c_summary = stats_data.get("company_summary", {})
    emp_details = stats_data.get("employee_details", {})
    
    # Lấy các bản ghi thô phát sinh trong tháng
    prefix = f"{year:04d}-{month:02d}-"
    monthly_records = [
        r for r in repo._attendance_records
        if str(r.get("date", "")).startswith(prefix)
    ]
    monthly_records.sort(key=lambda x: (str(x.get("date", "")), str(x.get("employee_id") or "")))

    # Tạo thư mục xuất file nếu chưa có
    os.makedirs(output_dir, exist_ok=True)
    filename = f"thong_ke_chuyen_can_thang_{month}_{year}.xlsx"
    file_path = os.path.join(output_dir, filename)

    # 2. Khởi tạo Workbook
    wb = openpyxl.Workbook()
    
    # Styles chuẩn cho báo cáo doanh nghiệp
    font_title = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="475569")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=11, color="1E293B")
    
    fill_header = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    fill_sub_header = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style="medium", color="1E40AF"))

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # ==========================================================
    # SHEET 1: TỔNG QUAN
    # ==========================================================
    ws1 = wb.active
    ws1.title = "Tổng Quan"
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = f"BÁO CÁO THỐNG KÊ CHUYÊN CẦN TOÀN CÔNG TY"
    ws1["A1"].font = font_title
    ws1["A2"] = f"Tháng {month} năm {year} | Xuất báo cáo ngày: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    ws1["A2"].font = font_subtitle

    # Header bảng tổng quan
    ws1["A4"] = "Chỉ số thống kê"
    ws1["B4"] = "Số lượng"
    ws1["C4"] = "Đơn vị tính"
    ws1["D4"] = "Tỷ lệ"

    for col in ["A", "B", "C", "D"]:
        cell = ws1[f"{col}4"]
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = cell_border

    tot_rec = c_summary.get("total_records", 0)
    pres = c_summary.get("present", 0)
    late = c_summary.get("late", 0)
    lv_app = c_summary.get("leave_approved", 0)
    abs_un = c_summary.get("absent_unexcused", 0)
    rate_pres = round((pres / tot_rec * 100), 1) if tot_rec > 0 else 0
    rate_late = round((late / tot_rec * 100), 1) if tot_rec > 0 else 0
    rate_lv = round((lv_app / tot_rec * 100), 1) if tot_rec > 0 else 0
    rate_abs = round((abs_un / tot_rec * 100), 1) if tot_rec > 0 else 0

    metrics = [
        ("Tổng số lượt chấm công", tot_rec, "Lượt", "100.0%"),
        ("Số lượt đúng giờ", pres, "Lượt", f"{rate_pres}%"),
        ("Số lượt đi muộn", late, "Lượt", f"{rate_late}%"),
        ("Số ngày vắng có phép", lv_app, "Ngày", f"{rate_lv}%"),
        ("Số ngày vắng không phép", abs_un, "Ngày", f"{rate_abs}%"),
        ("Số nhân sự phát sinh chấm công", len(emp_details), "Người", "-")
    ]

    for idx, (label, val, unit, pct) in enumerate(metrics, start=5):
        fill = fill_zebra if idx % 2 == 1 else fill_white
        ws1[f"A{idx}"] = label
        ws1[f"B{idx}"] = val
        ws1[f"C{idx}"] = unit
        ws1[f"D{idx}"] = pct

        ws1[f"A{idx}"].alignment = align_left
        ws1[f"B{idx}"].alignment = align_right
        ws1[f"C{idx}"].alignment = align_center
        ws1[f"D{idx}"].alignment = align_right

        for col in ["A", "B", "C", "D"]:
            cell = ws1[f"{col}{idx}"]
            cell.font = font_bold if idx == 5 else font_regular
            cell.fill = fill
            cell.border = cell_border

    # ==========================================================
    # SHEET 2: CHI TIẾT NHÂN VIÊN
    # ==========================================================
    ws2 = wb.create_sheet(title="Chi Tiết Nhân Viên")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = f"BẢNG THỐNG KÊ CHI TIẾT CHUYÊN CẦN TỪNG NHÂN VIÊN"
    ws2["A1"].font = font_title
    ws2["A2"] = f"Tháng {month} năm {year} | Tổng số nhân sự: {len(emp_details)}"
    ws2["A2"].font = font_subtitle

    headers_emp = [
        "STT", "Mã NV", "Họ và tên", "Phòng ban", "Chức vụ",
        "Tổng ngày", "Đúng giờ", "Đi muộn", "Vắng có phép", "Vắng không phép", "Tỷ lệ đúng giờ"
    ]

    for col_num, h_text in enumerate(headers_emp, start=1):
        c_letter = get_column_letter(col_num)
        cell = ws2[f"{c_letter}4"]
        cell.value = h_text
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = cell_border

    sorted_emps = sorted(emp_details.items(), key=lambda x: str(x[0]))
    row_idx = 5

    for idx, (emp_code, info) in enumerate(sorted_emps, start=1):
        s = info.get("summary", {})
        t_days = s.get("total_days", 0)
        p_days = s.get("present", 0)
        l_days = s.get("late", 0)
        lv_days = s.get("leave_approved", 0)
        ab_days = s.get("absent_unexcused", 0)
        emp_rate = f"{round((p_days / t_days * 100), 1)}%" if t_days > 0 else "0.0%"

        fill = fill_zebra if idx % 2 == 1 else fill_white
        row_data = [
            (idx, align_center),
            (emp_code, align_center),
            (info.get("name", emp_code), align_left),
            (info.get("department", "Công ty"), align_left),
            (info.get("position", "Nhân viên"), align_left),
            (t_days, align_right),
            (p_days, align_right),
            (l_days, align_right),
            (lv_days, align_right),
            (ab_days, align_right),
            (emp_rate, align_right)
        ]

        for c_num, (val, align) in enumerate(row_data, start=1):
            c_letter = get_column_letter(c_num)
            cell = ws2[f"{c_letter}{row_idx}"]
            cell.value = val
            cell.font = font_regular
            cell.fill = fill
            cell.alignment = align
            cell.border = cell_border
        
        row_idx += 1

    # Dòng tổng kết
    ws2[f"A{row_idx}"] = "TỔNG CỘNG"
    ws2[f"F{row_idx}"] = tot_rec
    ws2[f"G{row_idx}"] = pres
    ws2[f"H{row_idx}"] = late
    ws2[f"I{row_idx}"] = lv_app
    ws2[f"J{row_idx}"] = abs_un
    ws2[f"K{row_idx}"] = f"{rate_pres}%"

    ws2.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=5)
    for c_num in range(1, len(headers_emp) + 1):
        c_letter = get_column_letter(c_num)
        cell = ws2[f"{c_letter}{row_idx}"]
        cell.font = font_bold
        cell.fill = fill_sub_header
        cell.border = thick_bottom
        if c_num >= 6:
            cell.alignment = align_right
        else:
            cell.alignment = align_center

    # ==========================================================
    # SHEET 3: NHẬT KÝ CHẤM CÔNG (CHI TIẾT TỪNG BẢN GHI)
    # ==========================================================
    ws3 = wb.create_sheet(title="Nhật Ký Chấm Công")
    ws3.views.sheetView[0].showGridLines = True

    ws3["A1"] = f"NHẬT KÝ CHẤM CÔNG CHI TIẾT - THÁNG {month} NĂM {year}"
    ws3["A1"].font = font_title
    ws3["A2"] = f"Tổng số bản ghi phát sinh: {len(monthly_records)}"
    ws3["A2"].font = font_subtitle

    headers_log = [
        "STT", "Ngày", "Mã NV", "Họ và tên", "Phòng ban",
        "Giờ Check-in", "Giờ Check-out", "Trạng thái", "Ghi chú"
    ]

    for col_num, h_text in enumerate(headers_log, start=1):
        c_letter = get_column_letter(col_num)
        cell = ws3[f"{c_letter}4"]
        cell.value = h_text
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = cell_border

    st_vi_map = {
        "Present": "Đúng giờ",
        "Late": "Đi muộn",
        "Leave_Approved": "Vắng có phép",
        "Absent_Unexcused": "Vắng không phép"
    }

    for idx, r in enumerate(monthly_records, start=1):
        r_idx = idx + 4
        eid = str(r.get("employee_id") or r.get("employee_code") or "")
        emp_obj = repo._employees.get(eid, {})
        if not emp_obj and hasattr(repo, "_legacy_employees"):
            emp_obj = repo._legacy_employees.get(eid, {})

        emp_name = emp_obj.get("employee_name") or emp_obj.get("name") or eid
        emp_dept = emp_obj.get("department", "Công ty")
        st_raw = r.get("status", "")
        st_vi = st_vi_map.get(st_raw, st_raw)
        ci = r.get("check_in") or ""
        co = r.get("check_out") or ""
        note = r.get("notes") or ""

        fill = fill_zebra if idx % 2 == 1 else fill_white
        row_data = [
            (idx, align_center),
            (r.get("date", ""), align_center),
            (eid, align_center),
            (emp_name, align_left),
            (emp_dept, align_left),
            (ci, align_center),
            (co, align_center),
            (st_vi, align_center),
            (note, align_left)
        ]

        for c_num, (val, align) in enumerate(row_data, start=1):
            c_letter = get_column_letter(c_num)
            cell = ws3[f"{c_letter}{r_idx}"]
            cell.value = val
            cell.font = font_regular
            cell.fill = fill
            cell.alignment = align
            cell.border = cell_border

    # Tự động canh độ rộng cột cho cả 3 sheet
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if cell.row in [1, 2]:
                    continue  # Bỏ qua tiêu đề lớn để không làm giãn cột quá mức
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(file_path)

    return {
        "success": True,
        "filename": filename,
        "file_path": file_path,
        "month": month,
        "year": year,
        "total_records": tot_rec,
        "total_employees": len(emp_details)
    }
