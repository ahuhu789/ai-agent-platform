import os
import sys
import pytest

# Add root folder to sys.path
ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from agents.mcp_server import (
    get_attendance_history,
    get_employee_attendance,
    get_monthly_attendance_statistics,
    handle_rpc_request
)
from agents.mcp_client import MCPClient
from agents.attendance_agent import AttendanceAgent
from agents.contract import AgentRequest


# =========================================================================
# ATT-14: UNIT TESTS FOR 3 MCP TOOLS (Tốc độ siêu tốc < 0.1s, không cần Subprocess/Ollama)
# =========================================================================
class TestMCPToolsUnit:
    """Kiểm thử độc lập các hàm nghiệp vụ của 3 MCP Tool trực tiếp trong mcp_server.py"""

    def test_get_employee_attendance_success(self):
        # Case 1: Nhân viên hợp lệ 009 (Lê Hữu Thanh Vy)
        res1 = get_employee_attendance("009")
        assert res1["success"] is True
        assert res1["employee"]["employee_id"] == "009"
        assert res1["summary"]["total_days_tracked"] == 27
        assert len(res1["recent_records"]) <= 5

        # Case 2: Nhân viên hợp lệ 001 (Lê Thị Trà My)
        res2 = get_employee_attendance("001")
        assert res2["success"] is True
        assert "My" in res2["employee"]["name"]

    def test_get_employee_attendance_not_found(self):
        # Case nhân viên không tồn tại trong hệ thống
        res = get_employee_attendance("NON_EXISTENT_999")
        assert res["success"] is False
        assert "không tìm thấy" in res["error"].lower()

    def test_get_employee_attendance_zero_records(self):
        # Case ATT-13: Nhân viên 000001 (Ngô Hoàng Tâm) tồn tại nhưng chưa có bản ghi nào
        res = get_employee_attendance("000001")
        assert res["success"] is True
        assert "Ngô Hoàng Tâm" in res["employee"]["name"]
        assert res["summary"]["total_days_tracked"] == 0
        assert "chưa có bản ghi" in res.get("message", "").lower()

    def test_get_attendance_history_success(self):
        # Lịch sử trong khoảng ngày hợp lệ (009 có 9 bản ghi trong 01-15/08/2026)
        res = get_attendance_history("009", "2026-08-01", "2026-08-15")
        assert res["success"] is True
        assert res["employee"]["employee_id"] == "009"
        assert res["summary"]["total_days"] == 9
        assert len(res["records"]) == 9

    def test_get_attendance_history_reversed_dates(self):
        # Case ATT-12: start_date > end_date (Defense-in-depth)
        res = get_attendance_history("009", "2026-08-20", "2026-08-05")
        assert res["success"] is False
        assert "không hợp lệ" in res["error"].lower()
        assert "ngày bắt đầu" in res["error"].lower()

    def test_get_attendance_history_invalid_format(self):
        # Ngày sai định dạng hoặc ngày không tồn tại
        res = get_attendance_history("009", "2026-08-35", "2026-08-40")
        assert res["success"] is False
        assert "không hợp lệ" in res["error"].lower()

    def test_get_attendance_history_zero_records(self):
        # Nhân viên 000001 không có bản ghi trong khoảng ngày
        res = get_attendance_history("000001", "2026-08-01", "2026-08-15")
        assert res["success"] is True
        assert res["summary"]["total_days"] == 0
        assert len(res["records"]) == 0

    def test_get_monthly_statistics_success(self):
        # Tháng 9/2026 có 70 bản ghi
        res = get_monthly_attendance_statistics(9, 2026)
        assert res["success"] is True
        assert res["company_summary"]["total_records"] >= 50
        assert "001" in res["employee_details"]
        assert "009" in res["employee_details"]
        assert "010" in res["employee_details"]

    def test_get_monthly_statistics_no_data(self):
        # Tháng hợp lệ nhưng không có dữ liệu (tháng 10/2026)
        res = get_monthly_attendance_statistics(10, 2026)
        assert res["success"] is True
        assert "không có dữ liệu" in res.get("message", "").lower()

    def test_get_monthly_statistics_invalid_month(self):
        # Case ATT-11: Tháng ngoài khoảng 1-12
        res1 = get_monthly_attendance_statistics(13, 2026)
        assert res1["success"] is False
        assert "tháng không hợp lệ" in res1["error"].lower()

        res0 = get_monthly_attendance_statistics(0, 2026)
        assert res0["success"] is False
        assert "tháng không hợp lệ" in res0["error"].lower()

    def test_get_monthly_statistics_invalid_year(self):
        # Case ATT-11: Năm ngoài khoảng 2000-2030
        res_past = get_monthly_attendance_statistics(8, 1990)
        assert res_past["success"] is False
        assert "năm không hợp lệ" in res_past["error"].lower()

        res_future = get_monthly_attendance_statistics(8, 2035)
        assert res_future["success"] is False
        assert "năm không hợp lệ" in res_future["error"].lower()


# =========================================================================
# ATT-15 & ATT-19: SUBPROCESS LIFECYCLE & PROTOCOL TESTS
# =========================================================================
class TestMCPClientSubprocess:
    """Kiểm thử vòng đời tiến trình con MCPClient và chuẩn giao thức JSON-RPC"""

    def test_mcp_client_lifecycle_and_call(self):
        client = MCPClient()
        try:
            client.start()
            assert client.process is not None
            assert client.process.poll() is None

            # Gọi tool qua JSON-RPC stdio
            data = client.call_tool("get_employee_attendance", {"employee_id": "009"})
            assert data["success"] is True
            assert data["employee"]["employee_id"] == "009"
        finally:
            client.stop()
            assert client.process is None

    def test_mcp_client_list_tools(self):
        # Case ATT-19: list_tools qua JSON-RPC method 'tools/list'
        client = MCPClient()
        try:
            tools = client.list_tools()
            assert isinstance(tools, list)
            assert len(tools) >= 3
            tool_names = [t["name"] for t in tools]
            assert "get_attendance_history" in tool_names
            assert "get_employee_attendance" in tool_names
            assert "get_monthly_attendance_statistics" in tool_names
        finally:
            client.stop()

    def test_mcp_client_process_recovery(self):
        # Case ATT-15: Kiểm tra khả năng tự phục hồi khi subprocess bị kill đột ngột
        client = MCPClient()
        try:
            client.start()
            # Bắt chước sự cố: kill subprocess
            client.process.kill()
            client.process.wait()

            # Khi gọi tool kế tiếp, client phải tự khởi động lại process và trả về kết quả
            res = client.call_tool("get_employee_attendance", {"employee_id": "001"})
            assert res["success"] is True
            assert res["employee"]["employee_id"] == "001"
        finally:
            client.stop()


@pytest.fixture(scope="module")
def agent():
    ag = AttendanceAgent()
    yield ag
    ag.mcp_client.stop()

# =========================================================================
# AGENT VALIDATION LOGIC TESTS (ATT-11, ATT-12, ATT-13)
# =========================================================================
class TestAttendanceAgentValidation:
    """Kiểm thử các chốt chặn nghiệp vụ ở tầng Agent trước khi gọi tool"""

    def test_missing_employee_id(self, agent):
        req = AgentRequest(message="Xem chuyên cần của nhân viên")
        res = agent.handle(req)
        assert res.success is True
        assert "cung cấp mã nhân viên" in res.data.lower()

    def test_reversed_dates_rejected_by_agent(self, agent):
        # Case ATT-12
        req = AgentRequest(message="Cho tôi xem lịch sử chuyên cần của nhân viên 009 từ ngày 20/08/2026 đến 01/08/2026")
        res = agent.handle(req)
        assert res.success is True
        assert "không hợp lệ" in res.data.lower()
        assert "ngày bắt đầu" in res.data.lower()

    def test_invalid_month_rejected_by_agent(self, agent):
        # Case ATT-11: Tháng 13
        req = AgentRequest(message="Thống kê chuyên cần tháng 13 năm 2026")
        res = agent.handle(req)
        assert res.success is True
        assert "tháng 13 không hợp lệ" in res.data.lower()

    def test_invalid_year_rejected_by_agent(self, agent):
        # Case ATT-11: Năm 1990
        req = AgentRequest(message="Thống kê chuyên cần tháng 8 năm 1990")
        res = agent.handle(req)
        assert res.success is True
        # Agent từ chối hoặc hướng dẫn lại
        assert ("năm 1990 không hợp lệ" in res.data.lower()) or ("chuyên cần & chấm công" in res.data.lower()) or ("trợ lý ai tas" in res.data.lower())

    def test_unrelated_question(self, agent):
        req = AgentRequest(message="Thời tiết Hà Nội hôm nay thế nào?")
        res = agent.handle(req)
        assert res.success is True
        assert "chuyên cần & chấm công" in res.data.lower()


# =========================================================================
# ATT-17 & INTEGRATION TESTS (Ollama & End-to-End API)
# =========================================================================
class TestAttendanceIntegration:
    """Kiểm thử tích hợp đầu-cuối với LLM và FastAPI API endpoint"""

    def test_intent_extraction_use_case_1(self, agent):
        req = AgentRequest(message="Cho tôi xem lịch sử chuyên cần của nhân viên 009 từ ngày 01/08/2026 đến 15/08/2026")
        analysis = agent._extract_intent_and_params(req.message)
        assert analysis["intent"] == "attendance_history"
        assert analysis["employee_id"] == "009"
        assert analysis["from_date"] == "2026-08-01"
        assert analysis["to_date"] == "2026-08-15"

    def test_intent_extraction_use_case_2(self, agent):
        req = AgentRequest(message="Xem thông tin chuyên cần của nhân viên 001")
        analysis = agent._extract_intent_and_params(req.message)
        assert analysis["intent"] == "employee_attendance"
        assert analysis["employee_id"] == "001"

    def test_intent_extraction_use_case_3(self, agent):
        req = AgentRequest(message="Thống kê chuyên cần tháng 9 năm 2026")
        analysis = agent._extract_intent_and_params(req.message)
        assert analysis["intent"] == "monthly_attendance_statistics"
        assert int(analysis["month"]) == 9
        assert int(analysis["year"]) == 2026

    def test_edge_case_zero_attendance_end_to_end(self, agent):
        # Case ATT-13 qua Agent: Nhân viên 000001 (Ngô Hoàng Tâm) chưa có bản ghi nào
        req = AgentRequest(message="Xem thông tin chuyên cần của nhân viên 000001")
        res = agent.handle(req)
        assert res.success is True
        assert "ngô hoàng tâm" in res.data.lower() or "000001" in res.data
        assert "chưa có bản ghi chấm công" in res.data.lower() or "0" in res.data

    @pytest.mark.skipif(
        not os.path.exists(os.path.join(ROOT_DIR, "server.py")),
        reason="FastAPI server.py not present in team repository"
    )
    def test_api_chat_route_end_to_end(self):
        # Lazy import server.app ở đây để không làm chậm toàn bộ test suite nếu chỉ chạy unit test
        from fastapi.testclient import TestClient
        from server import app

        client = TestClient(app)
        try:
            response = client.post("/api/chat", json={
                "question": "Xem chuyên cần của nhân viên 001",
                "session_id": "test_attendance_api_session"
            })
            assert response.status_code == 200
            data = response.json()
            assert "answer" in data
            assert data["grounded"] is True
            assert len(data["sources"]) > 0
            assert "attendance.json" in data["sources"][0]["filename"]
        finally:
            client.delete("/api/sessions/test_attendance_api_session")

    def test_targeted_late_question(self, agent):
        # Case người dùng hỏi trọng tâm: Nhân viên 010 (Đào Minh Thuận) đi trễ bao nhiêu lần? (6 lần)
        req = AgentRequest(message="Nhân viên 010 đi trễ bao nhiêu lần?")
        res = agent.handle(req)
        assert res.success is True
        assert "6" in res.data or "đi trễ" in res.data or "muộn" in res.data
        assert "chức vụ:" not in res.data.lower()
        assert "phòng ban:" not in res.data.lower()

    def test_targeted_zero_late_question(self, agent):
        # Case người dùng hỏi nhân viên không đi trễ lần nào (014: Trần Đặng Ngọc Viên)
        req = AgentRequest(message="Nhân viên 014 đi trễ mấy lần?")
        res = agent.handle(req)
        assert res.success is True
        assert ("0 lần" in res.data) or ("không đi trễ" in res.data) or ("100%" in res.data) or ("0" in res.data)
        assert "chức vụ:" not in res.data.lower()

    def test_targeted_absence_days_without_checkin(self, agent):
        # Case người dùng hỏi nhân viên 009 có vắng ngày nào không -> Hiển thị số ngày không check-in có mặt đi làm
        req = AgentRequest(message="Nhân viên 009 có vắng ngày nào không?")
        res = agent.handle(req)
        assert res.success is True
        assert "009" in res.data or "lê hữu thanh vy" in res.data.lower()
        assert "không check-in" in res.data.lower()
        assert "9" in res.data
        assert "2026-09-01" in res.data
        assert "2026-09-12" in res.data

    def test_general_employee_attendance(self, agent):
        # Case người dùng hỏi tổng quan chuyên cần -> Trả về cấu trúc đầy đủ
        req = AgentRequest(message="Xem thông tin chuyên cần nhân viên 009")
        res = agent.handle(req)
        assert res.success is True
        assert "thông tin về nhân viên" in res.data.lower()
        assert "thống kê tổng quan:" in res.data.lower()
        assert "009" in res.data or "lê hữu thanh vy" in res.data.lower()

    def test_employee_list_query(self, agent):
        # Case tra cứu danh sách nhân viên
        req = AgentRequest(message="Xem danh sách nhân viên")
        res = agent.handle(req)
        assert res.success is True
        assert res.metadata["intent"] == "employee_list"
        assert "danh sách nhân viên" in res.data.lower()
        assert any(k in res.data.lower() for k in ["000001", "009", "001", "010", "ngô hoàng tâm", "lê hữu thanh vy"])

    def test_daily_attendance_absent_today(self, agent):
        # Case tra cứu nhân viên vắng mặt
        req = AgentRequest(message="Xem danh sách các nhân viên vắng mặt ngày 03/09/2026")
        res = agent.handle(req)
        assert res.success is True
        assert res.metadata["intent"] == "daily_attendance"
        assert "vắng" in res.data.lower() or "chưa ghi nhận" in res.data.lower()

    def test_daily_attendance_working_today(self, agent):
        # Case tra cứu nhân viên đi làm ngày 03/09/2026
        req = AgentRequest(message="Xem nhân viên đi làm ngày 03/09/2026")
        res = agent.handle(req)
        assert res.success is True
        assert res.metadata["intent"] == "daily_attendance"
        assert "đi làm" in res.data.lower() or "có mặt" in res.data.lower()
        assert any(k in res.data.lower() for k in ["001", "014", "015", "007", "lê thị trà my", "trần đặng ngọc viên"])

    def test_daily_attendance_vietnamese_date_format(self, agent):
        # Case người dùng hỏi "xem danh sách nhân viên đi làm ngày 18 tháng 9" -> Nhận diện đúng daily_attendance ngày 2026-09-18
        req = AgentRequest(message="xem danh sách nhân viên đi làm ngày 18 tháng 9")
        res = agent.handle(req)
        assert res.success is True
        assert res.metadata["intent"] == "daily_attendance"
        assert res.metadata["extracted_params"]["date"] == "2026-09-18"
        assert "2026-09-18" in res.data
        assert "0 nhân viên" in res.data or "chưa có" in res.data.lower()

    def test_daily_attendance_working_today_dynamic(self, agent):
        # Case tra cứu nhân viên đi làm hôm nay -> tự động lấy ngày hôm nay thực tế (datetime.now)
        from datetime import datetime
        today_str = datetime.now().strftime("%Y-%m-%d")
        req = AgentRequest(message="danh sách nhân viên có mặt hôm nay")
        res = agent.handle(req)
        assert res.success is True
        assert res.metadata["intent"] == "daily_attendance"
        assert today_str in res.data
        assert res.metadata["extracted_params"]["date"] == today_str

    def test_daily_attendance_absent_today_dynamic(self, agent):
        # Case tra cứu nhân viên vắng mặt hôm nay -> tự động lấy ngày hôm nay thực tế (datetime.now)
        from datetime import datetime
        today_str = datetime.now().strftime("%Y-%m-%d")
        req = AgentRequest(message="ai vắng mặt hôm nay?")
        res = agent.handle(req)
        assert res.success is True
        assert res.metadata["intent"] == "daily_attendance"
        assert today_str in res.data
        assert res.metadata["extracted_params"]["date"] == today_str

    def test_daily_attendance_specific_date_absent(self, agent):
        # Case tra cứu vắng mặt vào ngày cụ thể
        req = AgentRequest(message="Ai vắng mặt ngày 03/09/2026?")
        res = agent.handle(req)
        assert res.success is True
        assert res.metadata["intent"] == "daily_attendance"
        assert "vắng" in res.data.lower() or "chưa ghi nhận" in res.data.lower() or "danh sách" in res.data.lower()

    def test_export_monthly_attendance_excel_agent(self, agent):
        # Case người dùng yêu cầu xuất file excel thống kê chuyên cần tháng 7/2026
        import os
        import openpyxl
        req = AgentRequest(message="Xuất file excel thống kê tháng 7 năm 2026")
        res = agent.handle(req)
        assert res.success is True
        assert res.metadata["intent"] == "export_monthly_excel"
        assert "thong_ke_chuyen_can_thang_7_2026.xlsx" in res.data
        assert "/api/export/thong_ke_chuyen_can_thang_7_2026.xlsx" in res.data

        file_path = os.path.join("exports", "thong_ke_chuyen_can_thang_7_2026.xlsx")
        assert os.path.exists(file_path)
        wb = openpyxl.load_workbook(file_path)
        assert "Tổng Quan" in wb.sheetnames
        assert "Chi Tiết Nhân Viên" in wb.sheetnames
        assert "Nhật Ký Chấm Công" in wb.sheetnames

    @pytest.mark.skipif(
        not os.path.exists(os.path.join(ROOT_DIR, "server.py")),
        reason="FastAPI server.py not present in team repository"
    )
    def test_export_monthly_attendance_api_download(self):
        # Case kiểm tra endpoint API tải file Excel
        from fastapi.testclient import TestClient
        from server import app
        client = TestClient(app)
        response = client.get("/api/export/thong_ke_chuyen_can_thang_7_2026.xlsx")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers.get("content-type", "")
